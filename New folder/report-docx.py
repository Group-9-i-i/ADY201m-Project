from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Khởi tạo file Word
doc = Document()

# --- TIÊU ĐỀ ---
heading = doc.add_heading('BÁO CÁO XỬ LÝ DỮ LIỆU DỰ ÁN', 0)
heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Dự án: Dự báo Năng suất Cây trồng Karnataka')
doc.add_paragraph('Ngày báo cáo: 02/02/2026')
doc.add_paragraph('-' * 50)

# --- 1. TỔNG QUAN ---
doc.add_heading('1. Tổng quan', level=1)
doc.add_paragraph(
    'Mục tiêu báo cáo là tổng hợp quá trình làm sạch và chuẩn hóa dữ liệu (Data Cleaning Pipeline) '
    'để tạo ra bộ dữ liệu Master Dataset phục vụ huấn luyện mô hình AI.'
)

# --- 2. NGUỒN DỮ LIỆU ---
doc.add_heading('2. Nguồn dữ liệu đầu vào', level=1)
p = doc.add_paragraph()
p.add_run('• Dữ liệu Năng suất: ').bold = True
p.add_run('data_enriched_npk.csv (Diện tích, Sản lượng, Loại đất)\n')
p.add_run('• Dữ liệu Thời tiết: ').bold = True
p.add_run('Karnataka_Weather_Wind_2004_2025_Full.csv (Mưa, Nhiệt, Gió theo ngày)\n')
p.add_run('• Dữ liệu Viễn thám: ').bold = True
p.add_run('data_season_with_ndvi_fixed.csv (Chỉ số NDVI)')

# --- 3. QUY TRÌNH TÍCH HỢP ---
doc.add_heading('3. Quy trình Tích hợp Dữ liệu', level=1)
doc.add_paragraph(
    'Thực hiện gộp dữ liệu (Merging) dựa trên khóa chính là "District" và "Year". '
    'Dữ liệu thời tiết được tổng hợp (Aggregation) từ mức Ngày sang mức Năm:'
)
items = [
    'Lượng mưa: Tính TỔNG cả năm.',
    'Nhiệt độ & Gió: Tính TRUNG BÌNH cả năm.',
    'Chiến lược gộp: Left Join (Giữ nguyên dữ liệu cây trồng).'
]
for item in items:
    doc.add_paragraph(item, style='List Bullet')

# --- 4. CÁC BƯỚC LÀM SẠCH ---
doc.add_heading('4. Chi tiết 6 Bước Làm sạch', level=1)

steps = [
    ('Bước 1: Lọc cột & Đổi tên', 'Chuẩn hóa tên cột (Yield, Price) và loại bỏ cột thừa (lat, lon).'),
    ('Bước 2: Xử lý Trùng lặp', 'Xóa bỏ các dòng dữ liệu giống hệt nhau.'),
    ('Bước 3: Xử lý Dữ liệu thiếu', 'Điền Median cho biến số và Mode cho biến phân loại.'),
    ('Bước 4: Chuẩn hóa Định dạng', 'Viết hoa chữ cái đầu (Title Case), xóa khoảng trắng thừa.'),
    ('Bước 5: Nhất quán Dữ liệu', 'Sửa lỗi chính tả tên Quận (Mangalore -> Dakshina Kannada, Karnartakaka -> Karnataka).'),
    ('Bước 6: Khử nhiễu (Outliers)', 'Loại bỏ dữ liệu lỗi (Yield <= 0, Rainfall > 20,000mm).')
]

for step, desc in steps:
    p = doc.add_paragraph()
    p.add_run(f'{step}: ').bold = True
    p.add_run(desc)

# --- 5. KẾT QUẢ ---
doc.add_heading('5. Kết quả', level=1)
doc.add_paragraph('File thành phẩm: Cleaned_Master_Dataset_For_Training.csv')
doc.add_paragraph('Trạng thái: Đã sạch 100%, sẵn sàng cho Training Model.')

# Lưu file
file_name = 'Bao_Cao_Lam_Sach_Du_Lieu.docx'
doc.save(file_name)
print(f"Đã tạo xong file báo cáo: {file_name}")